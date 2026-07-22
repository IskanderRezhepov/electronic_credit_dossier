from pathlib import Path

from openpyxl import load_workbook

from app.services.document_reader import PageContent, ReadDocument
from app.services.exporter import save_excel
from app.services.table_extractor import extract_tables


def make_document(text: str, filename: str = "contract.pdf"):
    return ReadDocument(
        filename=filename,
        page_count=1,
        source_type="pdf",
        pages=[PageContent(1, text, "digital", len(text), 0.99, {"digital": text})],
    )


def test_extracts_equipment_type_model_quantity_vin_and_price():
    document = make_document(
        "СПЕЦИФИКАЦИЯ\n"
        "Самосвал марки SHACMAN X3000. Количество: 1 шт. "
        "VIN LZGJL4V44PX123456. Цена за единицу 45 000 000,00 тенге. "
        "Общая стоимость 45 000 000,00 тенге."
    )
    tables = extract_tables(document, "purchase_contract")
    equipment = next(table for table in tables if table["name"] == "asset_vin_rows")
    row = equipment["rows"][0]
    assert row["equipment_type"] == "Самосвал"
    assert row["quantity"] == 1
    assert row["vin"] == "LZGJL4V44PX123456"
    assert row["unit_price_kzt"] == 45000000.0
    assert row["total_amount_kzt"] == 45000000.0
    assert equipment["summary"]["total_quantity"] == 1


def test_counts_unique_vins_as_equipment_quantity():
    document = make_document(
        "ПРЕДМЕТ ЛИЗИНГА\n"
        "Погрузчик XCMG LW300FN VIN LZ0CC4W31P0123456\n"
        "Погрузчик XCMG LW300FN VIN LZ0CC4W31P0654321"
    )
    equipment = next(
        table for table in extract_tables(document, "lease_contract")
        if table["name"] == "asset_vin_rows"
    )
    assert equipment["summary"]["total_quantity"] == 2
    assert equipment["summary"]["unique_vin_count"] == 2
    assert equipment["summary"]["equipment_by_type"]["Погрузчик"] == 2


def test_does_not_invent_unit_price_from_unlabelled_amounts():
    document = make_document(
        "СПЕЦИФИКАЦИЯ Самосвал VIN LZGJL4V44PX123456 "
        "15 000 000,00 45 000 000,00"
    )
    equipment = next(
        table for table in extract_tables(document, "purchase_contract")
        if table["name"] == "asset_vin_rows"
    )
    row = equipment["rows"][0]
    assert row["unit_price_kzt"] is None


def test_excel_opens_on_summary_and_fields_are_sorted_by_page(tmp_path):
    path = tmp_path / "result.xlsx"
    data = {
        "documents": [{
            "document_type": "purchase_contract",
            "document_type_label_ru": "Договор купли-продажи",
            "used_ocr": False,
            "warnings": [],
            "fields": [
                {
                    "label_ru": "Поле со страницы 2", "value": "B", "page": 2,
                    "extraction_method": "digital", "confidence": 0.9,
                    "status": "extracted", "quote": None,
                },
                {
                    "label_ru": "Поле со страницы 1", "value": "A", "page": 1,
                    "extraction_method": "digital", "confidence": 0.9,
                    "status": "extracted", "quote": None,
                },
            ],
            "tables": [],
        }],
        "dossier": {"counts": {}, "checks": []},
    }
    save_excel(data, path)
    workbook = load_workbook(path)
    assert workbook.active.title == "Сводка"
    document_sheet = workbook["Договор купли-продажи"]
    assert document_sheet["A2"].value == "Поле со страницы 1"
    assert document_sheet["A3"].value == "Поле со страницы 2"
